"""Tests for bashbasicsbyvk_polyglot (the Polyglot Clone Finder).   python3 -m unittest discover -s tests/polyglot -v

The tool is installed flat (every file side by side in one bin/ folder), so the folders that hold its modules are put on sys.path here.
The launcher chdir()s into the folder it scans, so every call goes through main() below, which restores the working directory and
never installs or writes the xlsx; the tests that need those call the launcher directly.
"""
import atexit
import contextlib
import importlib.machinery
import importlib.util
import io
import json
import os
import shutil
import sys
import tempfile
import time
import unittest
from pathlib import Path

import numpy as np

HERE = Path(__file__).resolve().parent
SCRIPT = HERE.parents[1] / "script" / "file"
TOOL = SCRIPT / "infinity" / "tools" / "polyglot"
sys.path[:0] = [str(TOOL), str(SCRIPT / "engine"), str(SCRIPT / "infinity" / "tools" / "js_audit")]

import bashbasicsbyvk_pipboot as pipboot                                                            # noqa: E402
import bashbasicsbyvk_polyglot_bootstrap as bootstrap                                               # noqa: E402
import bashbasicsbyvk_polyglot_extractor as ex                                                      # noqa: E402
from bashbasicsbyvk_polyglot_constants import REPORT_NAME, WARNING                                  # noqa: E402
from bashbasicsbyvk_polyglot_crawler import crawl                                                   # noqa: E402
from bashbasicsbyvk_polyglot_extractor import FileInfo, FunctionBlock, TreeSitterBackend, extract_file  # noqa: E402
from bashbasicsbyvk_polyglot_languages import LANGUAGES                                             # noqa: E402
from bashbasicsbyvk_polyglot_metrics import compute_metrics                                         # noqa: E402
from bashbasicsbyvk_polyglot_normalize import canon_text, regex_tokens                              # noqa: E402
from bashbasicsbyvk_polyglot_pipeline import Config, run                                            # noqa: E402
from bashbasicsbyvk_polyglot_similarity import build_clusters, find_similar_pairs                   # noqa: E402
from bashbasicsbyvk_polyglot_ui import ProgressUI                                                   # noqa: E402
from bashbasicsbyvk_polyglot_verify import PairVerifier, dice_lcs, lcs_len, verify_pairs            # noqa: E402
from bashbasicsbyvk_polyglot_similarity import build_clusters as _bc                                # noqa: E402,F401

_loader = importlib.machinery.SourceFileLoader("bashbasicsbyvk_polyglot_launcher", str(TOOL / "bashbasicsbyvk_polyglot"))
launcher = importlib.util.module_from_spec(importlib.util.spec_from_loader(_loader.name, _loader))
_loader.exec_module(launcher)

_TMP = tempfile.mkdtemp(prefix="polyglot_test_")
atexit.register(shutil.rmtree, _TMP, True)
SAMPLE = Path(_TMP) / "sample_project"                      # a copy: the launcher writes its reports into the scanned folder
shutil.copytree(HERE / "sample_project", SAMPLE)
QUIET = ProgressUI(quiet=True)


def main(argv):
    cwd = os.getcwd()
    try:
        return launcher.main(list(argv) + ["--no-install", "--no-xlsx"])
    finally:
        os.chdir(cwd)


def fx(lang, src, name="f.x"):
    return ex.extract_file(name, src, LANGUAGES[lang], None)[0]


QUIET = ProgressUI(quiet=True)


def fx(lang, src, name="f.x"):
    return ex.extract_file(name, src, LANGUAGES[lang], None)[0]


class CrawlerTests(unittest.TestCase):
    def test_ignores_dirs_minified_and_unknown(self):
        found = {p.relative_to(SAMPLE).as_posix(): l for p, l in crawl(SAMPLE)}
        self.assertIn("src/finance.py", found)
        self.assertEqual(found["src/tax.go"], "go")
        self.assertEqual(found["src/slugify.ts"], "typescript")
        self.assertFalse(any("node_modules" in k for k in found))
        self.assertNotIn("src/bundle.min.js", found)

    def test_exclude_glob_and_lang_filter(self):
        got = {p.name for p, _ in crawl(SAMPLE, excludes=["*.go"], languages=["go", "python"])}
        self.assertNotIn("tax.go", got)
        self.assertIn("finance.py", got)
        self.assertNotIn("billing.js", got)


class NormalizeTests(unittest.TestCase):
    def test_identifier_split_and_cross_language_ops(self):
        self.assertEqual(canon_text("calculateTax"), ["calculate", "tax"])
        self.assertEqual(canon_text("compute_tax_rate"), ["compute", "tax", "rate"])
        self.assertEqual(canon_text("==="), ["=="])
        self.assertEqual(canon_text("None"), canon_text("null"))
        self.assertEqual(canon_text("let"), [])
        self.assertEqual(canon_text("{"), [])

    def test_comments_and_literals(self):
        toks = regex_tokens('x = "a}b"  # trailing comment with }\ny = 42', "hash")
        self.assertEqual(toks, ["x", "=", "<str>", "y", "=", "<num>"])
        toks = regex_tokens("a = 1; // note\n/* block */ b = 2.5", "c")
        self.assertEqual(toks, ["a", "=", "<num>", "b", "=", "<num>"])


class RegexExtractorTests(unittest.TestCase):
    def test_python_multiline_signature_nested_and_methods(self):
        src = ("class K:\n    def m(self,\n          a):\n        return a\n\n"
               "    def n(self):\n        def inner():\n            pass\n        return inner\nx = 1\n")
        got = [(f.name, f.start_line, f.end_line) for f in fx("python", src)]
        self.assertEqual(got, [("m", 2, 4), ("n", 6, 9)])      # nested def is not a separate block

    def test_javascript_braces_in_strings_and_comments(self):
        src = ('function a(x) {\n  if (x) { return "}"; } // }\n  return 1;\n}\n'
               "const b = (p) => {\n  return p;\n};\nclass C {\n  run(q) {\n    return q;\n  }\n}\n")
        got = {f.name: (f.start_line, f.end_line) for f in fx("javascript", src)}
        self.assertEqual(got, {"a": (1, 4), "b": (5, 7), "run": (9, 11)})

    def test_java_go_rust_ruby(self):
        java = "class T {\n  @Override\n  public double calc(double a,\n      double b) throws Exception {\n    return a;\n  }\n}\n"
        self.assertEqual([(f.name, f.start_line, f.end_line) for f in fx("java", java)], [("calc", 3, 6)])
        go = "package x\nfunc (s *S) Do(a interface{}) (int, error) {\n\treturn 1, nil\n}\nfunc Other() {\n}\n"
        self.assertEqual([f.name for f in fx("go", go)], ["Do", "Other"])
        rust = "trait T { fn decl(&self); }\nfn real(a: i32) -> i32 {\n    a + 1\n}\n"
        self.assertEqual([f.name for f in fx("rust", rust)], ["real"])
        ruby = "class A\n  def go(x)\n    if x\n      1\n    end\n  end\nend\n"
        self.assertEqual([(f.name, f.start_line, f.end_line) for f in fx("ruby", ruby)], [("go", 2, 6)])

    def test_c_and_cpp(self):
        c = "static int add(int a, int b) {\n  if (a) { return b; }\n  return a + b;\n}\n"
        self.assertEqual([f.name for f in fx("c", c)], ["add"])
        cpp = "int Foo::bar(int x) const {\n  return x;\n}\n"
        self.assertEqual([f.name for f in fx("cpp", cpp)], ["Foo::bar"])


class FakeNode:
    """Minimal stand-in for tree_sitter.Node, to exercise the AST code path offline."""
    def __init__(self, type_, text="", children=(), fields=None, line=0, end_line=0, start_byte=0):
        self.type, self._t, self.children = type_, text.encode(), list(children)
        self.fields, self.parent = fields or {}, None
        self.start_point, self.end_point = (line, 0), (end_line, 0)
        self.start_byte, self.end_byte = start_byte, start_byte + len(self._t)
        for c in self.children:
            c.parent = self

    @property
    def text(self): return self._t
    @property
    def child_count(self): return len(self.children)
    def child_by_field_name(self, n): return self.fields.get(n)


class AstPathTests(unittest.TestCase):
    def test_ast_extraction_skips_types_comments_and_name(self):
        name = FakeNode("identifier", "calculateTax", start_byte=4)
        body = [FakeNode("comment", "# hi"), FakeNode("type", "float"),
                FakeNode("string", "'s'", [FakeNode("string_content", "s")]),
                FakeNode("integer", "10"), FakeNode("identifier", "taxableAmount", start_byte=40),
                FakeNode("return", "return"), FakeNode("true", "True")]
        fn = FakeNode("function_definition", "", [FakeNode("def", "def"), name] + body, {"name": name}, 3, 9)
        root = FakeNode("module", "", [fn])
        blocks = ex._extract_ast(root, LANGUAGES["python"], "a.py", use_names=False)
        self.assertEqual(len(blocks), 1)
        b = blocks[0]
        self.assertEqual((b.name, b.start_line, b.end_line, b.backend), ("calculateTax", 4, 10, "tree-sitter"))
        self.assertEqual(b.tokens, ["<str>", "<num>", "taxable", "amount", "return", "true"])
        names_on = ex._extract_ast(root, LANGUAGES["python"], "a.py", use_names=True)[0].tokens
        self.assertEqual(names_on[:2], ["calculate", "tax"])


class MetricsAndClusterTests(unittest.TestCase):
    def test_overlapping_ranges_are_not_double_counted(self):
        prefix = list(range(0, 21))                         # 20 non-blank lines
        files = {"a.py": FileInfo("a.py", "python", 20, prefix)}
        funcs = [FunctionBlock("a.py", "f", "python", 1, 10, loc=10), FunctionBlock("a.py", "g", "python", 5, 15, loc=11)]
        from bashbasicsbyvk_polyglot_similarity import Cluster
        m = compute_metrics(files, funcs, [Cluster(1, [0, 1], [(0, 1, .9)])])
        self.assertEqual(m["duplicate_loc"], 15)
        self.assertEqual(m["duplication_pct"], 75.0)

    def test_pairs_and_components(self):
        v = np.array([[1, 0, 0], [1, .01, 0], [0, 1, 0], [0, 1, .01], [0, 0, 1]], dtype=np.float32)
        v /= np.linalg.norm(v, axis=1, keepdims=True)
        pairs, backend = find_similar_pairs(v, 0.95)
        self.assertEqual({(i, j) for i, j, _ in pairs}, {(0, 1), (2, 3)})
        cl = build_clusters(v, pairs)
        self.assertEqual(sorted(c.members for c in cl), [[0, 1], [2, 3]])
        self.assertTrue(all(c.avg_similarity > 0.99 for c in cl))


class PipelineTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.res = run(Config(str(SAMPLE)), QUIET)

    def _cluster_langs(self):
        return [sorted({self.res.funcs[m].language for m in c.members}) for c in self.res.clusters]

    def test_counts(self):
        self.assertEqual(len(self.res.files), 7)
        self.assertEqual(self.res.functions_extracted, 14)

    def test_cross_language_tax_cluster(self):
        tax = [c for c in self.res.clusters
               if {"calculate_tax", "computeTaxRate", "calculateTax", "CalcTax"} <= {self.res.funcs[m].name for m in c.members}]
        self.assertEqual(len(tax), 1)
        self.assertEqual(sorted(self._cluster_langs()[self.res.clusters.index(tax[0])]), ["go", "java", "javascript", "python"])

    def test_fibonacci_and_renamed_python_clusters(self):
        names = [{self.res.funcs[m].name for m in c.members} for c in self.res.clusters]
        self.assertIn({"fib", "fibonacci", "Fibonacci"}, names)
        self.assertIn({"reverse_words", "flip_sentence"}, names)

    def test_no_false_positives_and_vendored_copy_ignored(self):
        in_cluster = {self.res.funcs[m].name for c in self.res.clusters for m in c.members}
        for unrelated in ("slugify", "load_settings", "clamp", "renderInvoiceRow", "describe"):
            self.assertNotIn(unrelated, in_cluster)
        self.assertTrue(all("node_modules" not in f.file for f in self.res.funcs))

    def test_metrics_formula(self):
        m = self.res.metrics
        dup = sum(f.loc for c in self.res.clusters for f in (self.res.funcs[i] for i in c.members))
        self.assertEqual(m["duplicate_loc"], dup)
        self.assertEqual(m["total_loc"], sum(i.loc for i in self.res.files.values()))
        self.assertAlmostEqual(m["duplication_pct"], round(100 * dup / m["total_loc"], 2))
        self.assertTrue(all(0 <= r["duplication_pct"] <= 100 for r in m["files"]))

    def test_cross_language_only(self):
        res = run(Config(str(SAMPLE), cross_language_only=True), QUIET)
        names = [{res.funcs[m].name for m in c.members} for c in res.clusters]
        self.assertNotIn({"reverse_words", "flip_sentence"}, names)
        self.assertEqual(len(res.clusters), 2)

    def test_empty_and_missing_dirs(self):
        with tempfile.TemporaryDirectory() as d:
            res = run(Config(d), QUIET)
            self.assertEqual((len(res.files), res.metrics["duplication_pct"]), (0, 0.0))
        from bashbasicsbyvk_polyglot_pipeline import AnalysisError
        with self.assertRaises(AnalysisError):
            run(Config("/definitely/not/here"), QUIET)


class PythonOneLinerTests(unittest.TestCase):
    def test_one_liner_defs_do_not_swallow_following_code(self):
        src = ("class A:\n    def cursor(self): return (self.r, self.c)\n    def _lf(self):\n        self.r += 1\n"
               "    def wheel(self, up, t=1): self.send(b'x', t)\n    def close(self):\n        pass\n"
               "def one(a):  return a and (a +\n    1)\ndef after(x): pass\n")
        got = [(f.name, f.start_line, f.end_line) for f in fx("python", src)]
        self.assertEqual(got, [("cursor", 2, 2), ("_lf", 3, 4), ("wheel", 5, 5), ("close", 6, 7),
                               ("one", 8, 9), ("after", 10, 10)])

    def test_annotated_multiline_header_and_docstring(self):
        src = 'def f(a,\n      b: dict[str, int] = {}) -> "x":\n    """def fake(): ..."""\n    return a\n'
        self.assertEqual([(f.name, f.start_line, f.end_line) for f in fx("python", src)], [("f", 1, 4)])


def _blk(name, toks, file="a.py", start=1, end=10):
    b = FunctionBlock(file, name, "python", start, end, toks)
    return b


class VerificationTests(unittest.TestCase):
    A = "x = foo ( a ) if x > 0 else bar ( b ) for i in range ( n ) : total += i return total".split()

    def _verify(self, funcs, pairs, thr=0.55, ratio=0.4):
        for i, f in enumerate(funcs):
            f.id = i
        v = PairVerifier(funcs, thr)
        return verify_pairs(funcs, pairs, thr, ratio, v), v

    def test_identical_and_renamed_pass_reordered_bag_fails(self):
        same = _blk("a", self.A * 3, "a.py")
        renamed = _blk("b", [("v" + t if t.isalpha() and t not in ("if", "else", "for", "in", "range", "return") else t)
                             for t in self.A * 3], "b.py")
        shuffled = _blk("c", sorted(self.A * 3), "c.py")                  # same bag of tokens, no structure
        (pairs, st), v = self._verify([same, renamed, shuffled], [(0, 1, .9), (0, 2, .9)])
        self.assertEqual([(i, j) for i, j, _ in pairs], [(0, 1)])
        self.assertEqual(st["rejected_verify"], 1)

    def test_size_ratio_and_overlap_guards(self):
        big, small = _blk("big", self.A * 10, end=100), _blk("small", self.A * 2, file="b.py")
        inner = _blk("inner", self.A * 10, start=5, end=40)       # inside `big`'s line range, same file
        (pairs, st), _ = self._verify([big, small, inner], [(0, 1, .95), (0, 2, .95)])
        self.assertEqual(pairs, [])
        self.assertEqual((st["rejected_size"], st["rejected_overlap"]), (1, 1))

    def test_average_linkage_pruning_breaks_chains(self):
        # a~b and b~c are strong, a!~c: single linkage glues a,b,c; average linkage must drop one.
        a = _blk("a", "if x : return y".split() * 6, "a.py")
        c = _blk("c", "for i in range ( n ) : total += i".split() * 6, "c.py")
        b = _blk("b", ("if x : return y".split() * 3) + ("for i in range ( n ) : total += i".split() * 3), "b.py")
        funcs = [a, b, c]
        for i, f in enumerate(funcs):
            f.id = i
        v = PairVerifier(funcs, 0.0)
        x = np.eye(3, dtype=np.float32)
        cl = build_clusters(x, [(0, 1, .9), (1, 2, .9)], sim_fn=v.sim, prune_below=0.6)
        self.assertTrue(all(len(c_.members) == 2 for c_ in cl))
        self.assertLessEqual(len(cl), 1)

    def test_pipeline_reports_verification_and_can_be_disabled(self):
        on = run(Config(str(SAMPLE)), QUIET)
        off = run(Config(str(SAMPLE), verify=False), QUIET)
        self.assertIsNotNone(on.verify_stats)
        self.assertIsNone(off.verify_stats)
        self.assertGreaterEqual(sum(len(c.members) for c in off.clusters),
                                sum(len(c.members) for c in on.clusters))


class BashTests(unittest.TestCase):
    SRC = ("#!/usr/bin/env bash\n"
           "greet() {\n  local n=\"$1\"   # } in a comment\n  echo \"hi } ${n}\" ; echo '}' \\}\n}\n"
           "function build_all {\n  cat <<EOF\nunbalanced { and 'quote\nEOF\n  [[ $# -gt 0 ]] && echo \"${#}\"\n}\n"
           "nl_style()\n{\n  return 0\n}\n"
           "outer() {\n  inner() { :; }\n  inner\n}\n"
           "x=$(( 1 << 3 ))\n")

    def test_extraction_forms_strings_heredocs_and_comments(self):
        got = [(f.name, f.start_line, f.end_line) for f in fx("bash", self.SRC)]
        self.assertEqual(got, [("greet", 2, 5), ("build_all", 6, 11), ("nl_style", 12, 15), ("outer", 16, 19)])

    def test_tokens_variables_flags_and_comments(self):
        toks = regex_tokens('local n="$1"  # note\nif [[ -n "$x" && $# -gt 0 ]]; then echo ${name} --force; fi', "bash")
        self.assertNotIn("note", toks)
        self.assertIn("name", toks)
        self.assertIn("--force", toks)
        self.assertIn("-n", toks)
        self.assertNotIn("local", toks)

    def test_shebang_detection_for_extensionless_scripts(self):
        with tempfile.TemporaryDirectory() as d:
            for name, body in (("tool", "#!/bin/bash\nf() {\n:\n}\n"), ("pytool", "#!/usr/bin/env python3\nprint(1)\n"),
                               ("notes", "plain text\n"), ("LICENSE", "MIT\n"), ("run.sh", "f() { :; }\n")):
                Path(d, name).write_text(body)
            found = {p.name: l for p, l in crawl(Path(d))}
            self.assertEqual(found, {"tool": "bash", "pytool": "python", "run.sh": "bash"})

    def test_bash_clone_detected_and_siblings_not(self):
        body = ('{n}() {{\n  local f="$1" k v\n  [ -f "$f" ] || return 1\n  while IFS= read -r line; do\n'
                '    case "$line" in\n      \\#*) continue ;;\n      *=*) k="${{line%%=*}}"; v="${{line#*=}}" ;;\n    esac\n'
                '    printf -v "$k" %s "$v"\n  done < "$f"\n  return 0\n}}\n')
        other = ('stop_spinner() {\n  [ -n "$PID" ] || return 0\n  kill -KILL "$PID" 2>/dev/null\n  wait "$PID" 2>/dev/null\n'
                 '  PID=""\n  printf "\\033[?25h" >"$TTY"\n  rm -f "$LOCK"\n}}\n')
        with tempfile.TemporaryDirectory() as d:
            Path(d, "a.sh").write_text(body.format(n="read_conf") + "\n" + other)
            Path(d, "b.sh").write_text(body.format(n="load_settings").replace("local f=", "local cfg=").replace('"$f"', '"$cfg"'))
            res = run(Config(d, min_lines=5), QUIET)
        names = [{res.funcs[m].name for m in c.members} for c in res.clusters]
        self.assertIn({"read_conf", "load_settings"}, names)
        self.assertFalse(any("stop_spinner" in n for n in names))


def _dp_lcs(a, b):
    prev = [0] * (len(b) + 1)
    for x in a:
        cur = [0]
        for j, y in enumerate(b):
            cur.append(prev[j] + 1 if x == y else max(prev[j + 1], cur[j]))
        prev = cur
    return prev[-1]


class NoModelTests(unittest.TestCase):
    def test_no_ml_dependencies_are_imported(self):
        run(Config(str(SAMPLE)), QUIET)
        self.assertFalse({"torch", "transformers", "sklearn", "faiss"} & set(sys.modules))

    def test_codebert_options_are_gone(self):
        for flag in ("--embedder", "--model", "--device", "--no-center"):
            with self.assertRaises(SystemExit):
                main([str(SAMPLE), flag, "x", "-q"])


class LcsTests(unittest.TestCase):
    def test_bit_parallel_lcs_equals_dynamic_programming(self):
        import random
        rnd = random.Random(7)
        for _ in range(300):
            a = [rnd.choice("abcde") for _ in range(rnd.randint(0, 50))]
            b = [rnd.choice("abcdef") for _ in range(rnd.randint(0, 50))]
            self.assertEqual(lcs_len(a, b), _dp_lcs(a, b))

    def test_dice_properties(self):
        t = "a b c d e f".split()
        self.assertEqual(dice_lcs(t, t), 1.0)
        self.assertEqual(dice_lcs(t, list(reversed(t))), dice_lcs(list(reversed(t)), t))
        self.assertEqual(dice_lcs(t, []), 0.0)
        self.assertEqual(dice_lcs(t, "x y z".split()), 0.0)


class PythonAstBackendTests(unittest.TestCase):
    SRC = ("class A:\n    def cursor(self): return (self.r, self.c)\n    @staticmethod\n    def multi(a,\n"
           "              b: int = 3) -> dict:\n        \"\"\"doc\"\"\"\n        s = f\"{a}-{b!r}\"  # c\n        def inner(): pass\n"
           "        return {k: 0 for k in range(10)}\nasync def top(x):\n    return await x\n")

    def test_stdlib_ast_extents_without_tree_sitter(self):
        blocks, used = extract_file("f.py", self.SRC, LANGUAGES["python"], TreeSitterBackend())
        if used == "tree-sitter":
            self.skipTest("tree-sitter-python installed; the stdlib path is not taken")
        self.assertEqual(used, "python-ast")
        self.assertEqual([(b.name, b.start_line, b.end_line) for b in blocks],
                         [("cursor", 2, 2), ("multi", 4, 9), ("top", 10, 11)])      # closure `inner` stays in `multi`
        self.assertEqual(blocks[1].tokens.count("<str>"), 2)                       # docstring + the whole f-string

    def test_falls_back_to_regex_on_syntax_error(self):
        blocks, used = extract_file("f.py", "def ok(a):\n    return a\ndef broken(:\n", LANGUAGES["python"], TreeSitterBackend())
        self.assertIn(used, ("regex", "tree-sitter"))
        self.assertIn("ok", [b.name for b in blocks])


class FastPathAndSplitTests(unittest.TestCase):
    def test_exact_clones_are_linked_and_scored_100(self):
        body = "def f(a, b):\n    t = 0\n    for i in range(a):\n        if i % 3 == 0:\n            t += i * b\n    return t\n"
        with tempfile.TemporaryDirectory() as d:
            for n in ("a.py", "b.py", "c.py"):
                Path(d, n).write_text(body)
            res = run(Config(d, min_lines=4, min_tokens=10), QUIET)
        self.assertEqual(len(res.clusters), 1)
        self.assertEqual(len(res.clusters[0].members), 3)
        self.assertAlmostEqual(res.clusters[0].avg_similarity, 1.0)
        self.assertGreaterEqual(res.verify_stats["exact_clone_links"], 2)

    def test_oversized_chain_is_split_not_kept(self):
        # 60 functions in 20 tight triples, glued into ONE connected component by weak (0.62) bridge edges
        n = 60
        x = np.eye(n, dtype=np.float32)
        pairs = []
        sim = {}
        for g in range(20):
            a, b, c = 3 * g, 3 * g + 1, 3 * g + 2
            for u, v in ((a, b), (b, c), (a, c)):
                pairs.append((u, v, 0.95))
                sim[(u, v)] = 0.95
            if g:
                pairs.append((a - 1, a, 0.62))
                sim[(a - 1, a)] = 0.62
        fn = lambda i, j: sim.get((min(i, j), max(i, j)), 0.0)          # noqa: E731
        cl = build_clusters(x, pairs, sim_fn=fn, prune_below=0.6, max_refine=40)
        self.assertTrue(cl)
        self.assertLessEqual(max(len(c.members) for c in cl), 40)
        self.assertTrue(all(c.avg_similarity >= 0.6 for c in cl))

    def test_blind_skeleton_match_rejected_but_renamed_copy_accepted(self):
        body = "for i in range ( n ) : if x : return y".split()
        a = _blk("a", body + ["p", "q", "p", "r", "q", "p"] * 6, "a.py")      # same skeleton, but identifiers bind
        b = _blk("b", body + ["zz%d" % i for i in range(36)], "b.py")        # in a completely different pattern
        real = "t = foo ( a ) for i in range ( n ) : if i > k : t += bar ( i , a ) return t".split()
        ren = [("q" + w if w.isalpha() and w not in ("for", "in", "range", "if", "return") else w) for w in real]
        c, d = _blk("c", real * 2, "c.py"), _blk("d", ren * 2, "d.py")        # consistent renaming of every identifier
        v = PairVerifier([a, b, c, d], 0.5)
        self.assertLess(v.sim(0, 1), 0.5)
        self.assertAlmostEqual(v.sim(2, 3), 0.9)

    def test_min_confidence_filters_clusters(self):
        allc = run(Config(str(SAMPLE)), QUIET).clusters
        high = run(Config(str(SAMPLE), min_confidence="high"), QUIET).clusters
        self.assertLessEqual(len(high), len(allc))
        self.assertTrue(all(c.avg_similarity >= 0.80 for c in high))


class CliTests(unittest.TestCase):
    def test_json_report_schema_and_exit_codes(self):
        with tempfile.TemporaryDirectory() as d:
            out = os.path.join(d, "report.json")
            self.assertEqual(main([str(SAMPLE), "-o", out, "--plain", "-q"]), 0)
            rep = json.loads(Path(out).read_text())
            s = rep["summary"]
            for k in ("total_files_scanned", "total_functions_extracted", "overall_duplication_pct", "duplicate_clusters"):
                self.assertIn(k, s)
            self.assertEqual(s["total_files_scanned"], 7)
            c0 = rep["clusters"][0]
            self.assertTrue({"members", "matches", "avg_similarity_pct", "cross_language"} <= set(c0))
            self.assertTrue({"file", "function", "language", "start_line", "end_line"} <= set(c0["members"][0]))
            self.assertTrue(rep["files"] and "duplication_pct" in rep["files"][0])
            self.assertEqual(main([str(SAMPLE), "-o", out, "-q", "--fail-above", "10"]), 1)
            self.assertEqual(main([str(SAMPLE), "-o", out, "-q", "--fail-above", "99"]), 0)
            self.assertEqual(main(["/nope", "-o", out, "-q"]), 2)
            self.assertEqual(main([str(SAMPLE), "-o", out, "-q", "--verify-threshold", "2"]), 2)
            self.assertEqual(main([str(SAMPLE), "-o", out, "-q", "--no-verify"]), 0)
            rep = json.loads(Path(out).read_text())
            self.assertFalse(rep["summary"]["verification"]["enabled"])


class LauncherTests(unittest.TestCase):
    def _copy(self):
        d = Path(tempfile.mkdtemp(dir=_TMP)) / "p"
        shutil.copytree(SAMPLE, d, ignore=shutil.ignore_patterns("Polyglot Report.*"))
        return d

    def test_runs_inside_the_scanned_folder_and_writes_both_reports_there(self):
        d, cwd = self._copy(), os.getcwd()
        try:
            self.assertEqual(launcher.main([str(d), "--no-install", "--plain", "-q"]), 0)
            self.assertEqual(Path(os.getcwd()).resolve(), d.resolve())           # cd'd into it, like the audits
        finally:
            os.chdir(cwd)
        self.assertTrue((d / f"{REPORT_NAME}.json").is_file())
        self.assertTrue((d / f"{REPORT_NAME}.xlsx").is_file())

    def test_warning_is_in_console_json_and_xlsx(self):
        d, cwd = self._copy(), os.getcwd()
        out = io.StringIO()
        try:
            with contextlib.redirect_stdout(out):
                launcher.main([str(d), "--no-install", "--plain"])
        finally:
            os.chdir(cwd)
        self.assertIn(WARNING, out.getvalue())
        self.assertIn("NOT a 100% reliable", out.getvalue())
        self.assertEqual(json.loads((d / f"{REPORT_NAME}.json").read_text())["disclaimer"], WARNING)
        import openpyxl
        wb = openpyxl.load_workbook(d / f"{REPORT_NAME}.xlsx")
        self.assertEqual(wb.sheetnames, ["Summary", "Clusters", "Files"])
        for ws in wb:
            self.assertIn("NOT a 100% reliable", ws["A2"].value)
        self.assertGreater(wb["Clusters"].max_row, 4)

    def test_folder_without_sources_is_not_an_error(self):
        d, cwd = Path(tempfile.mkdtemp(dir=_TMP)), os.getcwd()
        out = io.StringIO()
        try:
            with contextlib.redirect_stdout(out):
                self.assertEqual(launcher.main([str(d), "--no-install"]), 0)
        finally:
            os.chdir(cwd)
        self.assertIn("No supported source files", out.getvalue())


class BootstrapTests(unittest.TestCase):
    """The pip calls are faked: nothing is ever installed by the tests."""

    def setUp(self):
        d = Path(tempfile.mkdtemp(dir=_TMP))
        self._saved = (bootstrap.HOME_DIR, bootstrap.SENTINEL_DIR, bootstrap.LOG_FILE, pipboot.have, pipboot.run_pip)
        bootstrap.HOME_DIR, bootstrap.SENTINEL_DIR, bootstrap.LOG_FILE = d, d / "s", d / "log"
        self.pip_calls, self.installed, self.failing = [], set(), set()
        pipboot.have = lambda module, dist=None, min_version=None: module in self.installed
        pipboot.run_pip = self._fake_pip

    def tearDown(self):
        bootstrap.HOME_DIR, bootstrap.SENTINEL_DIR, bootstrap.LOG_FILE, pipboot.have, pipboot.run_pip = self._saved

    def _fake_pip(self, args, log, timeout):
        spec = args[-1]
        self.pip_calls.append(spec)
        if spec in self.failing:
            return 1
        self.installed.add(spec.split(">=")[0].replace("-", "_"))
        return 0

    def test_grammar_spec_is_the_module_name_with_dashes(self):
        self.assertEqual(bootstrap.grammar_spec("tree_sitter_c_sharp"), "tree-sitter-c-sharp>=0.23")

    def test_installs_core_then_only_the_grammars_of_languages_found(self):
        self.installed = {"numpy", "rich", "openpyxl"}
        missing = bootstrap.bootstrap(["python", "typescript", "tsx"], "auto", lambda m: None)
        self.assertEqual(missing, [])
        self.assertEqual(self.pip_calls, ["tree-sitter>=0.23", "tree-sitter-python>=0.23", "tree-sitter-typescript>=0.23"])

    def test_failed_core_skips_grammars_and_is_not_retried_within_a_day(self):
        self.installed = {"numpy", "rich", "openpyxl"}
        self.failing = {"tree-sitter>=0.23"}
        bootstrap.bootstrap(["python", "go"], "auto", lambda m: None)
        self.assertEqual(self.pip_calls.count("tree-sitter>=0.23"), 2)         # --user, then --user --break-system-packages: one attempt
        self.pip_calls.clear()
        bootstrap.bootstrap(["python", "go"], "auto", lambda m: None)
        self.assertEqual(self.pip_calls, [])                                    # the sentinel says "failed < 24h ago"
        old = time.time() - 2 * bootstrap.RETRY
        for s in bootstrap.SENTINEL_DIR.iterdir():
            os.utime(s, (old, old))
        bootstrap.bootstrap(["python", "go"], "auto", lambda m: None)
        self.assertIn("tree-sitter>=0.23", self.pip_calls)                      # the window has passed: tried again

    def test_regex_parser_never_installs_tree_sitter(self):
        self.installed = {"numpy", "rich", "openpyxl"}
        bootstrap.bootstrap(["python"], "regex", lambda m: None)
        self.assertEqual(self.pip_calls, [])

    def test_missing_numpy_is_reported_as_required(self):
        self.failing = {"numpy"}
        self.assertEqual(bootstrap.bootstrap(["python"], "regex", lambda m: None), ["numpy"])

    def test_installed_everything_means_no_pip_call(self):
        self.installed = {"numpy", "rich", "openpyxl", "tree_sitter", "tree_sitter_python"}
        bootstrap.bootstrap(["python"], "auto", lambda m: None)
        self.assertEqual(self.pip_calls, [])


class PipbootTests(unittest.TestCase):
    def test_have_checks_module_and_minimum_version(self):
        self.assertTrue(pipboot.have("numpy"))
        self.assertTrue(pipboot.have("numpy", "numpy", (0, 1)))
        self.assertFalse(pipboot.have("numpy", "numpy", (999,)))
        self.assertFalse(pipboot.have("no_such_module_xyz"))

    def test_retry_due(self):
        s = Path(tempfile.mkdtemp(dir=_TMP)) / "sentinel"
        self.assertTrue(pipboot.retry_due(s, 60))
        s.write_text("x")
        self.assertFalse(pipboot.retry_due(s, 60))
        os.utime(s, (time.time() - 120, time.time() - 120))
        self.assertTrue(pipboot.retry_due(s, 60))


if __name__ == "__main__":
    unittest.main(verbosity=2)
