"""bashbasicsbyvk_polyglot_xlsx - the formatted Excel report (same colours and font as the JavaScript audit workbook).

Three sheets: Summary (with the reliability warning), Clusters (one row per function, grouped by cluster), Files (per-file duplication).
Needs openpyxl; the launcher installs it and skips this report if that fails.
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from bashbasicsbyvk_js_audit_constants import (
    CLR_HEADER_BG, CLR_HEADER_FG, CLR_TITLE_BG, CLR_TITLE_FG, CLR_OK_BG, CLR_OK_FG, CLR_INFO_BG, CLR_INFO_FG,
    CLR_WARN_BG, CLR_WARN_FG, CLR_ROW_ODD, CLR_ROW_EVEN, FONT_NAME,
)
from bashbasicsbyvk_polyglot_constants import TOOL_NAME, VERSION, WARNING

BAND = {"high": (CLR_OK_BG, CLR_OK_FG), "medium": (CLR_INFO_BG, CLR_INFO_FG), "low": (CLR_WARN_BG, CLR_WARN_FG)}


def _put(ws, row, col, value, bold=False, bg=None, fg="000000", h="left"):
    cell = ws.cell(row, col, value)
    cell.font = Font(name=FONT_NAME, bold=bold, color=fg, size=10)
    cell.alignment = Alignment(horizontal=h, vertical="center", wrap_text=True)
    side = Side(style="thin", color="D0D0D0")
    cell.border = Border(left=side, right=side, top=side, bottom=side)
    if bg:
        cell.fill = PatternFill("solid", fgColor=bg)
    return cell


def _banner(ws, title, ncols):
    """Row 1 = title, row 2 = the reliability warning, both merged across the sheet."""
    for row, text, bg, fg, height in ((1, title, CLR_TITLE_BG, CLR_TITLE_FG, 24), (2, "⚠  " + WARNING, CLR_WARN_BG, CLR_WARN_FG, 44)):
        _put(ws, row, 1, text, bold=True, bg=bg, fg=fg, h="center")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
        ws.row_dimensions[row].height = height


def _sheet(wb, title, heading, headers, widths, rows):
    """rows = [(values, row_bg, {col_index: (bg, fg)})]; header on row 3, data from row 4."""
    ws = wb.create_sheet(title)
    _banner(ws, heading, len(headers))
    for c, name in enumerate(headers, 1):
        _put(ws, 3, c, name, bold=True, bg=CLR_HEADER_BG, fg=CLR_HEADER_FG, h="center")
        ws.column_dimensions[get_column_letter(c)].width = widths[c - 1]
    for r, (values, bg, special) in enumerate(rows, 4):
        for c, v in enumerate(values, 1):
            colours = special.get(c)
            _put(ws, r, c, v, bold=bool(colours), bg=colours[0] if colours else bg, fg=colours[1] if colours else "000000",
                 h="center" if isinstance(v, (int, float)) or colours else "left")
    ws.freeze_panes = ws.cell(4, 1)
    ws.sheet_view.showGridLines = False
    if rows:
        ws.auto_filter.ref = f"A3:{get_column_letter(len(headers))}{3 + len(rows)}"
    return ws


def write_xlsx(report, path):
    s = report["summary"]
    wb = Workbook()
    wb.remove(wb.active)

    pairs = [
        ("Tool", f"{TOOL_NAME} {VERSION}"), ("Scanned folder", report["root"]), ("Generated", report["generated_at"]),
        ("Files scanned", s["total_files_scanned"]), ("Functions extracted", s["total_functions_extracted"]),
        ("Functions analysed (>= min size)", s["functions_analyzed"]),
        ("Duplicate clusters", f"{s['duplicate_clusters']} ({s['cross_language_clusters']} cross-language)"),
        ("Functions in clusters", s["functions_in_duplicate_clusters"]),
        ("Total LOC / duplicate LOC", f"{s['total_loc']} / {s['duplicate_loc']}"),
        ("Duplication of the project (candidates)", f"{s['overall_duplication_pct']:.2f}%"),
        ("Redundant LOC (candidates)", f"{s['redundant_loc']} ({s['redundant_pct']:.2f}%)"),
        ("Parser per language", ", ".join(f"{k}: {v}" for k, v in sorted(s["parser_backend_by_language"].items())) or "-"),
        ("Elapsed seconds", s["elapsed_seconds"]),
    ] + [("Note", n) for n in report["notes"]]
    ws = wb.create_sheet("Summary")
    _banner(ws, f"{TOOL_NAME} Report -- Summary", 2)
    for r, (k, v) in enumerate(pairs, 3):
        _put(ws, r, 1, k, bold=True, bg=CLR_HEADER_BG, fg=CLR_HEADER_FG, h="center")
        _put(ws, r, 2, v, bg=CLR_ROW_ODD if r % 2 else CLR_ROW_EVEN)
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 110
    ws.sheet_view.showGridLines = False

    rows = []
    for n, c in enumerate(report["clusters"]):
        bg = CLR_ROW_ODD if n % 2 == 0 else CLR_ROW_EVEN
        for m in c["members"]:
            rows.append(([c["cluster_id"], c["confidence"].upper(), c["avg_similarity_pct"], c["size"], m["function"], m["file"],
                          f"{m['start_line']}-{m['end_line']}", m["language"], m["loc"], m["similarity_to_cluster_pct"]],
                         bg, {2: BAND[c["confidence"]]}))
    _sheet(wb, "Clusters", f"{TOOL_NAME} Report -- Candidate duplicate clusters",
           ["Cluster", "Confidence", "Avg Match %", "Size", "Function", "File", "Lines", "Language", "LOC", "Match to Cluster %"],
           [9, 13, 13, 8, 34, 52, 12, 12, 8, 18], rows)

    rows = [([f["path"], f["language"], f["loc"], f["functions_in_duplicate_clusters"], f["duplicate_loc"], f["duplication_pct"]],
             CLR_ROW_ODD if i % 2 == 0 else CLR_ROW_EVEN, {})
            for i, f in enumerate(sorted(report["files"], key=lambda r: -r["duplication_pct"]))]
    _sheet(wb, "Files", f"{TOOL_NAME} Report -- Duplication per file",
           ["File", "Language", "LOC", "Functions in Clusters", "Duplicate LOC", "Duplication %"], [60, 12, 8, 22, 16, 16], rows)
    wb.save(path)
